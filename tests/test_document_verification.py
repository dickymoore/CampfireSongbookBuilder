import json
import tempfile
import unittest
from pathlib import Path

from app.document_verification import (
    build_document_verification_record,
    evaluate_document_artifact,
    evaluate_markdown_artifact_neatness,
    load_document_verification,
    save_document_verification,
)
from tests.docx_stub import install_docx_stub

install_docx_stub()

from docx import Document


class TestDocumentVerification(unittest.TestCase):
    def _build_record(
        self,
        artifact_path,
        artifact_type,
        verification_status,
        verification_reasons=None,
        verified_at="2026-05-21T16:00:00+01:00",
    ):
        return build_document_verification_record(
            artifact_path=artifact_path,
            artifact_type=artifact_type,
            verification_status=verification_status,
            verification_reasons=verification_reasons,
            verified_at=verified_at,
        )

    def test_save_and_load_round_trip_for_multiple_artifacts(self):
        markdown_record = self._build_record(
            "data/output/Lyrics_Document.md",
            "markdown",
            "failed",
            verification_reasons=["placeholder reason supplied by later neatness checks"],
        )
        docx_record = self._build_record(
            "data/output/Lyrics_Document.docx",
            "docx",
            "passed",
            verification_reasons=[],
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            target_path = Path(tmp_dir) / "data" / "review" / "document_quality.json"

            save_document_verification(target_path, [markdown_record, docx_record])

            self.assertTrue(target_path.exists())
            self.assertTrue(target_path.parent.exists())

            state, errors = load_document_verification(target_path)

            self.assertEqual(errors, [])
            self.assertEqual(state["version"], 1)
            self.assertEqual(
                state["entries"]["data/output/Lyrics_Document.md"],
                markdown_record,
            )
            self.assertEqual(
                state["entries"]["data/output/Lyrics_Document.docx"],
                docx_record,
            )

    def test_load_missing_file_returns_empty_state(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            target_path = Path(tmp_dir) / "data" / "review" / "document_quality.json"

            state, errors = load_document_verification(target_path)

            self.assertEqual(errors, [])
            self.assertEqual(state["version"], 1)
            self.assertIsNone(state["updated_at"])
            self.assertEqual(state["entries"], {})

    def test_load_malformed_json_reports_validation_error(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            target_path = Path(tmp_dir) / "data" / "review" / "document_quality.json"
            target_path.parent.mkdir(parents=True, exist_ok=True)
            target_path.write_text("{not valid json", encoding="utf-8")

            state, errors = load_document_verification(target_path)

            self.assertEqual(state["entries"], {})
            self.assertEqual(len(errors), 1)
            self.assertEqual(errors[0]["file_path"], str(target_path))
            self.assertEqual(errors[0]["field"], "$")
            self.assertIn("JSON", errors[0]["reason"])

    def test_load_invalid_nested_entry_preserves_valid_records(self):
        valid_record = self._build_record(
            "data/output/Lyrics_Document.docx",
            "docx",
            "passed",
            verification_reasons=[],
        )

        payload = {
            "version": 1,
            "updated_at": "2026-05-21T16:00:00+01:00",
            "entries": {
                "data/output/Lyrics_Document.docx": valid_record,
                "data/output/Lyrics_Document.md": {
                    "artifact_path": "data/output/Lyrics_Document.md",
                    "artifact_type": "markdown",
                    "verification_status": "failed",
                    "verification_reasons": "not-a-list",
                    "verified_at": "2026-05-21T16:00:00+01:00",
                },
            },
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            target_path = Path(tmp_dir) / "data" / "review" / "document_quality.json"
            target_path.parent.mkdir(parents=True, exist_ok=True)
            target_path.write_text(json.dumps(payload), encoding="utf-8")

            state, errors = load_document_verification(target_path)

            self.assertEqual(
                state["entries"]["data/output/Lyrics_Document.docx"],
                valid_record,
            )
            self.assertNotIn("data/output/Lyrics_Document.md", state["entries"])
            self.assertEqual(len(errors), 1)
            self.assertEqual(
                errors[0]["field"],
                "entries['data/output/Lyrics_Document.md'].verification_reasons",
            )
            self.assertIn("must be a list", errors[0]["reason"])

    def test_load_entry_key_mismatch_reports_validation_error(self):
        payload = {
            "version": 1,
            "updated_at": "2026-05-21T16:00:00+01:00",
            "entries": {
                "data/output/Lyrics_Document.md": {
                    "artifact_path": "data/output/Other_Document.md",
                    "artifact_type": "markdown",
                    "verification_status": "failed",
                    "verification_reasons": ["placeholder reason supplied by later neatness checks"],
                    "verified_at": "2026-05-21T16:00:00+01:00",
                }
            },
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            target_path = Path(tmp_dir) / "data" / "review" / "document_quality.json"
            target_path.parent.mkdir(parents=True, exist_ok=True)
            target_path.write_text(json.dumps(payload), encoding="utf-8")

            state, errors = load_document_verification(target_path)

            self.assertEqual(state["entries"], {})
            self.assertEqual(len(errors), 1)
            self.assertEqual(
                errors[0]["field"],
                "entries['data/output/Lyrics_Document.md'].artifact_path",
            )
            self.assertIn("entry key must match artifact_path", errors[0]["reason"])

    def test_build_record_rejects_non_list_verification_reasons(self):
        with self.assertRaisesRegex(
            ValueError,
            "verification_reasons must be a list",
        ):
            build_document_verification_record(
                artifact_path="data/output/Lyrics_Document.md",
                artifact_type="markdown",
                verification_status="failed",
                verification_reasons="manual verification pending",
                verified_at="2026-05-21T16:00:00+01:00",
            )

    def test_distinguishes_near_collision_artifact_paths(self):
        primary_record = self._build_record(
            "data/output/Lyrics_Document.docx",
            "docx",
            "passed",
            verification_reasons=[],
        )
        secondary_record = self._build_record(
            "data/output/Lyrics_Document.md",
            "markdown",
            "failed",
            verification_reasons=["manual verification pending"],
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            target_path = Path(tmp_dir) / "data" / "review" / "nested" / "document_quality.json"

            save_document_verification(target_path, [primary_record, secondary_record])

            state, errors = load_document_verification(target_path)

            self.assertEqual(errors, [])
            self.assertEqual(len(state["entries"]), 2)
            self.assertIn("data/output/Lyrics_Document.docx", state["entries"])
            self.assertIn("data/output/Lyrics_Document.md", state["entries"])

    def test_evaluate_markdown_artifact_neatness_passes_clean_song_blocks(self):
        evaluation = evaluate_markdown_artifact_neatness(
            "\n".join(
                [
                    "# Trail Song by The Campfire Trio",
                    "",
                    "```text",
                    "First line",
                    "Second line",
                    "```",
                    "",
                    "# River Song by The Campfire Trio",
                    "",
                    "```text",
                    "Third line",
                    "Fourth line",
                    "```",
                    "",
                ]
            )
        )

        self.assertEqual(evaluation["verification_status"], "passed")
        self.assertEqual(
            evaluation["verification_reasons"],
            ["meets_neatness_thresholds"],
        )

    def test_evaluate_document_artifact_reports_failing_markdown_reasons(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            artifact_path = Path(tmp_dir) / "lyrics.md"
            artifact_path.write_text(
                "\n".join(
                    [
                        "# Trail Song by The Campfire Trio",
                        "",
                        "```text",
                        "Only line",
                        "```",
                        "",
                        "",
                        "",
                        "# River Song by The Campfire Trio",
                        "",
                        "```text",
                        "Solo",
                        "```",
                        "",
                    ]
                )
                + "\n",
                encoding="utf-8",
            )

            record = evaluate_document_artifact(artifact_path, artifact_type="markdown")

            self.assertEqual(record["verification_status"], "failed")
            self.assertIn("excessive_whitespace", record["verification_reasons"])
            self.assertIn("sparse_layout", record["verification_reasons"])
            self.assertIn("fragmented_song_blocks", record["verification_reasons"])

    def test_evaluate_document_artifact_reports_failing_docx_reasons(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            artifact_path = Path(tmp_dir) / "lyrics.docx"
            document = Document()
            document.add_heading("Trail Song by The Campfire Trio", level=1)
            document.add_paragraph("Only line")
            document.add_heading("River Song by The Campfire Trio", level=1)
            document.add_paragraph("Solo")
            document.save(artifact_path)

            record = evaluate_document_artifact(artifact_path, artifact_type="docx")

            self.assertEqual(record["verification_status"], "failed")
            self.assertIn("sparse_layout", record["verification_reasons"])
            self.assertIn("fragmented_song_blocks", record["verification_reasons"])

    def test_evaluate_document_artifact_does_not_treat_body_text_as_heading(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            artifact_path = Path(tmp_dir) / "lyrics.docx"
            document = Document()
            document.add_heading("Trail Song by The Campfire Trio", level=1)
            document.add_paragraph("Sung by the firelight\nSecond line")
            document.save(artifact_path)

            record = evaluate_document_artifact(artifact_path, artifact_type="docx")

            self.assertEqual(record["verification_status"], "passed")
            self.assertEqual(
                record["verification_reasons"],
                ["meets_neatness_thresholds"],
            )


if __name__ == "__main__":
    unittest.main()
