import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from app.content_models import build_favourite_record, build_review_decision, compute_content_hash
from app.document_verification import load_document_verification
from app.review_state import load_content_scores, save_review_decisions
from tests.docx_stub import install_docx_stub
from tests.pdf_fixtures import write_minimal_pdf

install_docx_stub()

try:
    from docx import Document
    from app.document_creation import create_document_from_cache
    DOCX_AVAILABLE = True
except ModuleNotFoundError:
    Document = None
    create_document_from_cache = None
    DOCX_AVAILABLE = False


@unittest.skipUnless(DOCX_AVAILABLE, "python-docx is not installed in this environment")
class TestDocumentCreation(unittest.TestCase):
    def _read_document_text(self, path):
        document = Document(path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)

    def _run_create_document(self, tmp_path, *args, **kwargs):
        import app.document_creation as document_creation

        original_document_quality_path = document_creation.DOCUMENT_QUALITY_PATH
        original_content_scores_path = document_creation.CONTENT_SCORES_PATH
        document_creation.DOCUMENT_QUALITY_PATH = (
            tmp_path / "data" / "review" / "document_quality.json"
        )
        document_creation.CONTENT_SCORES_PATH = tmp_path / "data" / "review" / "content_scores.json"
        try:
            return create_document_from_cache(*args, **kwargs)
        finally:
            document_creation.DOCUMENT_QUALITY_PATH = original_document_quality_path
            document_creation.CONTENT_SCORES_PATH = original_content_scores_path

    def test_create_document_from_cache_excludes_questionable_and_overlong_content_by_default(self):
        song_list = [
            {"Artist": "The Campfire Trio", "Title": "Trail Song"},
            {"Artist": "The Campfire Trio", "Title": "Long Song"},
            {"Artist": "The Campfire Trio", "Title": "Missing Song"},
        ]
        lyrics_cache = {
            "The Campfire Trio - Trail Song": "First line\nSecond line",
            "The Campfire Trio - Long Song": "\n".join(
                "Line {}".format(index) for index in range(5100)
            ),
            "The Campfire Trio - Missing Song": "Lyrics not found.",
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "lyrics.docx"

            report_data = self._run_create_document(
                tmp_path,
                song_list,
                lyrics_cache,
                {},
                lyrics_output=output_path,
            )

            text = self._read_document_text(output_path)
            markdown_text = output_path.with_suffix(".md").read_text(encoding="utf-8")
            report_entries = {entry["title"]: entry for entry in report_data["entries"]}

            self.assertIn("Trail Song by The Campfire Trio", text)
            self.assertNotIn("Missing Song by The Campfire Trio", text)
            self.assertNotIn("Long Song by The Campfire Trio", text)
            self.assertIn("Trail Song by The Campfire Trio", markdown_text)
            self.assertNotIn("Missing Song by The Campfire Trio", markdown_text)
            self.assertNotIn("Long Song by The Campfire Trio", markdown_text)
            self.assertTrue(report_entries["Trail Song"]["included"])
            self.assertFalse(report_entries["Missing Song"]["included"])
            self.assertFalse(report_entries["Long Song"]["included"])
            self.assertIn(
                "print_hostile_content",
                [signal["code"] for signal in report_entries["Long Song"]["signals"]],
            )
            self.assertEqual(
                report_entries["Long Song"]["reason"],
                "Questionable content is excluded by default.",
            )

    def test_create_document_from_cache_can_include_questionable_content_by_mode(self):
        song_list = [
            {"Artist": "The Campfire Trio", "Title": "Long Song"},
            {"Artist": "The Campfire Trio", "Title": "Missing Song"},
        ]
        lyrics_cache = {
            "The Campfire Trio - Long Song": "\n".join(
                "Line {}".format(index) for index in range(220)
            ),
            "The Campfire Trio - Missing Song": "Lyrics not found.",
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "lyrics.docx"

            report_data = self._run_create_document(
                tmp_path,
                song_list,
                lyrics_cache,
                {},
                lyrics_output=output_path,
                include_questionable=True,
            )

            text = self._read_document_text(output_path)
            report_entries = {entry["title"]: entry for entry in report_data["entries"]}

            self.assertIn("Long Song by The Campfire Trio", text)
            self.assertTrue(report_entries["Long Song"]["included"])
            self.assertEqual(
                report_entries["Long Song"]["reason"],
                "Questionable content included by generation mode.",
            )
            self.assertFalse(report_entries["Missing Song"]["included"])

    def test_create_document_from_cache_adds_contents_page_before_songs(self):
        song_list = [
            {"Artist": "The Campfire Trio", "Title": "Trail Song"},
            {"Artist": "The Campfire Trio", "Title": "River Song"},
        ]
        lyrics_cache = {
            "The Campfire Trio - Trail Song": "First line\nSecond line",
            "The Campfire Trio - River Song": "Third line\nFourth line",
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "lyrics.docx"

            self._run_create_document(
                tmp_path,
                song_list,
                lyrics_cache,
                {},
                lyrics_output=output_path,
            )

            document = Document(output_path)
            paragraph_texts = [paragraph.text for paragraph in document.paragraphs]

            self.assertEqual(paragraph_texts[0], "2 Campfire Songs | 2026-07-09")
            self.assertEqual(len(document.tables), 1)
            table = document.tables[0]
            self.assertEqual(table.cell(1, 0).text, "River Song - The Campfire Trio")
            self.assertEqual(table.cell(2, 0).text, "Trail Song - The Campfire Trio")
            self.assertIn("Trail Song by The Campfire Trio", paragraph_texts)

    def test_create_document_from_cache_persists_content_scores_for_lyrics_and_chords(self):
        song_list = [{"Artist": "The Campfire Trio", "Title": "Trail Song"}]
        lyrics_cache = {"The Campfire Trio - Trail Song": "First line\nSecond line"}
        chords_cache = {"The Campfire Trio - Trail Song": "[G]Trail song"}

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            lyrics_output_path = tmp_path / "lyrics.docx"
            chords_output_path = tmp_path / "chords.docx"

            report_data = self._run_create_document(
                tmp_path,
                song_list,
                lyrics_cache,
                chords_cache,
                lyrics_output=lyrics_output_path,
                chords_output=chords_output_path,
            )

            content_scores_path = tmp_path / "data" / "review" / "content_scores.json"
            state, errors = load_content_scores(content_scores_path)

            self.assertEqual(errors, [])
            self.assertIn("The Campfire Trio - Trail Song", state["entries"])
            self.assertIn("lyrics", state["entries"]["The Campfire Trio - Trail Song"])
            self.assertIn("chords", state["entries"]["The Campfire Trio - Trail Song"])

            lyric_score = state["entries"]["The Campfire Trio - Trail Song"]["lyrics"]
            chord_score = state["entries"]["The Campfire Trio - Trail Song"]["chords"]
            self.assertEqual(lyric_score["quality_score"], 100)
            self.assertEqual(lyric_score["quality_band"], "clean")
            self.assertEqual(chord_score["quality_score"], 100)
            self.assertEqual(chord_score["quality_band"], "clean")

            self.assertTrue(isinstance(report_data.get("content_scores"), list))
            self.assertEqual(
                {(record["song_key"], record["content_type"]) for record in report_data["content_scores"]},
                {("The Campfire Trio - Trail Song", "lyrics"), ("The Campfire Trio - Trail Song", "chords")},
            )

    def test_create_document_from_cache_reports_chords_wrap_risk_audit(self):
        song_list = [{"Artist": "The Campfire Trio", "Title": "Trail Song"}]
        chords_cache = {
            "The Campfire Trio - Trail Song": "X" * 61 + "\nShort line"
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "chords.docx"

            report_data = self._run_create_document(
                tmp_path,
                song_list,
                {},
                chords_cache,
                chords_output=output_path,
                include_questionable=True,
            )

            self.assertEqual(len(report_data["chords_layout_audit"]), 1)
            audit_entry = report_data["chords_layout_audit"][0]
            self.assertEqual(audit_entry["song_key"], "The Campfire Trio - Trail Song")
            self.assertEqual(audit_entry["font_name"], "Courier New")
            self.assertEqual(audit_entry["font_size"], 8)
            self.assertTrue(audit_entry["has_overlong_lines"])
            self.assertEqual(audit_entry["overlong_line_count"], 1)
            self.assertEqual(audit_entry["overlong_lines"][0]["line_number"], 1)

    def test_create_document_from_cache_includes_questionable_content_with_override(self):
        song_list = [
            {"Artist": "The Campfire Trio", "Title": "Trail Song"},
            {"Artist": "The Campfire Trio", "Title": "Override Song"},
        ]
        lyrics_cache = {
            "The Campfire Trio - Trail Song": "First line\nSecond line",
            "The Campfire Trio - Override Song": "Lyrics not found.",
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            review_decisions_path = tmp_path / "review_decisions.json"
            override_decision = build_review_decision(
                "The Campfire Trio - Override Song",
                "lyrics",
                compute_content_hash("Lyrics not found."),
                "override",
                reason="Approved for this print run.",
                decided_at="2026-05-19T15:00:34+01:00",
            )
            save_review_decisions(review_decisions_path, [override_decision])

            output_path = tmp_path / "lyrics.docx"

            import app.document_creation as document_creation

            original_quality_status_path = document_creation.QUALITY_STATUS_PATH
            original_review_decisions_path = document_creation.REVIEW_DECISIONS_PATH
            document_creation.QUALITY_STATUS_PATH = tmp_path / "missing_quality_status.json"
            document_creation.REVIEW_DECISIONS_PATH = review_decisions_path
            try:
                self._run_create_document(
                    tmp_path,
                    song_list,
                    lyrics_cache,
                    {},
                    lyrics_output=output_path,
                )
            finally:
                document_creation.QUALITY_STATUS_PATH = original_quality_status_path
                document_creation.REVIEW_DECISIONS_PATH = original_review_decisions_path

            text = self._read_document_text(output_path)

            self.assertIn("Trail Song by The Campfire Trio", text)
            self.assertIn("Override Song by The Campfire Trio", text)

    def test_create_document_from_cache_filters_selected_songs_in_selection_order(self):
        selection_records = [
            build_favourite_record("The Campfire Trio", "Trail Song"),
            build_favourite_record("The Campfire Trio", "Missing Song"),
            build_favourite_record("The Campfire Trio", "Override Song"),
        ]
        lyrics_cache = {
            "The Campfire Trio - Trail Song": "First line\nSecond line",
            "The Campfire Trio - Missing Song": "Lyrics not found.",
            "The Campfire Trio - Override Song": "Lyrics not found.",
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            review_decisions_path = tmp_path / "review_decisions.json"
            override_decision = build_review_decision(
                "The Campfire Trio - Override Song",
                "lyrics",
                compute_content_hash("Lyrics not found."),
                "override",
                reason="Approved for this selected book.",
                decided_at="2026-05-19T15:00:34+01:00",
            )
            save_review_decisions(review_decisions_path, [override_decision])

            output_path = tmp_path / "lyrics.docx"

            import app.document_creation as document_creation

            original_quality_status_path = document_creation.QUALITY_STATUS_PATH
            original_review_decisions_path = document_creation.REVIEW_DECISIONS_PATH
            document_creation.QUALITY_STATUS_PATH = tmp_path / "missing_quality_status.json"
            document_creation.REVIEW_DECISIONS_PATH = review_decisions_path
            try:
                report_data = self._run_create_document(
                    tmp_path,
                    [],
                    lyrics_cache,
                    {},
                    lyrics_output=output_path,
                    selection_records=selection_records,
                    report_source="generate_from_selection",
                )
            finally:
                document_creation.QUALITY_STATUS_PATH = original_quality_status_path
                document_creation.REVIEW_DECISIONS_PATH = original_review_decisions_path

            text = self._read_document_text(output_path)
            report_entries = {entry["title"]: entry for entry in report_data["entries"]}

            self.assertIn("Trail Song by The Campfire Trio", text)
            self.assertIn("Override Song by The Campfire Trio", text)
            self.assertLess(
                text.index("Trail Song by The Campfire Trio"),
                text.index("Override Song by The Campfire Trio"),
            )
            self.assertNotIn("Missing Song by The Campfire Trio", text)
            self.assertTrue(report_entries["Trail Song"]["included"])
            self.assertFalse(report_entries["Missing Song"]["included"])
            self.assertEqual(
                report_entries["Missing Song"]["reason"],
                "Questionable content is excluded by default.",
            )
            self.assertTrue(report_entries["Override Song"]["included"])
            self.assertEqual(report_entries["Override Song"]["decision_source"], "review_override")
            self.assertEqual(report_data["source"], "generate_from_selection")

    def test_create_document_from_cache_reports_selection_completeness_issues_before_output(self):
        selection_records = [
            build_favourite_record("The Campfire Trio", "Trail Song"),
            build_favourite_record("The Campfire Trio", "Missing Song"),
        ]
        lyrics_cache = {
            "The Campfire Trio - Trail Song": "First line\nSecond line",
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "lyrics.docx"

            report_data = self._run_create_document(
                tmp_path,
                [],
                lyrics_cache,
                {},
                lyrics_output=output_path,
                selection_records=selection_records,
                report_source="generate_from_selection",
            )

            text = self._read_document_text(output_path)
            report_entries = {entry["title"]: entry for entry in report_data["entries"]}

            self.assertIn("Trail Song by The Campfire Trio", text)
            self.assertNotIn("Missing Song by The Campfire Trio", text)
            self.assertTrue(report_entries["Trail Song"]["included"])
            self.assertFalse(report_entries["Missing Song"]["included"])
            self.assertEqual(
                report_entries["Missing Song"]["reason"],
                "Missing content is excluded by default.",
            )
            self.assertEqual(len(report_data["selection_issues"]), 1)
            self.assertEqual(report_data["selection_issues"][0]["song_key"], "The Campfire Trio - Missing Song")
            self.assertEqual(report_data["selection_issues"][0]["issue_type"], "missing_content")
            self.assertEqual(report_data["source"], "generate_from_selection")

    def test_create_document_from_cache_excludes_song_from_both_documents_when_one_side_is_missing(self):
        song_list = [{"Artist": "The Campfire Trio", "Title": "Half Missing Song"}]
        lyrics_cache = {"The Campfire Trio - Half Missing Song": "First line\nSecond line"}
        chords_cache = {"The Campfire Trio - Half Missing Song": "Chords not found."}

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            lyrics_output = tmp_path / "lyrics.docx"
            chords_output = tmp_path / "chords.docx"

            report_data = self._run_create_document(
                tmp_path,
                song_list,
                lyrics_cache,
                chords_cache,
                lyrics_output=lyrics_output,
                chords_output=chords_output,
            )

            lyrics_text = self._read_document_text(lyrics_output)
            chords_text = self._read_document_text(chords_output)
            report_entries = {
                (entry["title"], entry["content_type"]): entry
                for entry in report_data["entries"]
            }

            self.assertNotIn("Half Missing Song by The Campfire Trio", lyrics_text)
            self.assertNotIn("Half Missing Song by The Campfire Trio", chords_text)
            self.assertFalse(report_entries[("Half Missing Song", "lyrics")]["included"])
            self.assertFalse(report_entries[("Half Missing Song", "chords")]["included"])
            self.assertEqual(
                report_entries[("Half Missing Song", "lyrics")]["reason"],
                "Song is missing lyrics or chords and is excluded from both documents.",
            )
            self.assertEqual(
                report_entries[("Half Missing Song", "chords")]["reason"],
                "Song is missing lyrics or chords and is excluded from both documents.",
            )

    def test_create_document_from_cache_writes_pdf_when_converter_succeeds(self):
        song_list = [{"Artist": "The Campfire Trio", "Title": "Trail Song"}]
        lyrics_cache = {"The Campfire Trio - Trail Song": "First line\nSecond line"}

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "lyrics.docx"
            document_quality_path = tmp_path / "data" / "review" / "document_quality.json"

            def fake_convert(docx_path, pdf_path=None, runner=None, which=None):  # noqa: ARG001
                target_path = Path(docx_path).with_suffix(".pdf")
                write_minimal_pdf(target_path, "x" * 80)
                return target_path, None

            with patch("app.document_creation.convert_document_to_pdf", side_effect=fake_convert):
                report_data = self._run_create_document(
                    tmp_path,
                    song_list,
                    lyrics_cache,
                    {},
                    lyrics_output=output_path,
                    pdf_output=True,
                )

            self.assertTrue(output_path.exists())
            self.assertTrue(output_path.with_suffix(".md").exists())
            self.assertTrue(output_path.with_suffix(".pdf").exists())
            self.assertEqual(report_data["pdf_outputs"], [str(output_path.with_suffix(".pdf"))])
            self.assertEqual(report_data["pdf_errors"], [])
            state, errors = load_document_verification(document_quality_path)
            self.assertEqual(errors, [])
            pdf_records = [
                record
                for record in report_data["document_verification"]
                if record.get("artifact_type") == "pdf"
            ]
            self.assertEqual(len(pdf_records), 1)
            self.assertEqual(pdf_records[0]["artifact_path"], str(output_path.with_suffix(".pdf")))
            self.assertIn(str(output_path.with_suffix(".pdf")), state["entries"])

    def test_create_document_from_cache_preserves_artifacts_when_pdf_conversion_fails(self):
        song_list = [{"Artist": "The Campfire Trio", "Title": "Trail Song"}]
        lyrics_cache = {"The Campfire Trio - Trail Song": "First line\nSecond line"}

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "lyrics.docx"
            document_quality_path = tmp_path / "data" / "review" / "document_quality.json"

            with patch(
                "app.document_creation.convert_document_to_pdf",
                return_value=(None, "converter missing"),
            ):
                report_data = self._run_create_document(
                    tmp_path,
                    song_list,
                    lyrics_cache,
                    {},
                    lyrics_output=output_path,
                    pdf_output=True,
                )

            self.assertTrue(output_path.exists())
            self.assertTrue(output_path.with_suffix(".md").exists())
            self.assertFalse(output_path.with_suffix(".pdf").exists())
            self.assertEqual(report_data["pdf_outputs"], [])
            self.assertEqual(len(report_data["pdf_errors"]), 1)
            self.assertEqual(report_data["pdf_errors"][0]["reason"], "converter missing")
            state, errors = load_document_verification(document_quality_path)
            self.assertEqual(errors, [])
            self.assertNotIn(str(output_path.with_suffix(".pdf")), state["entries"])
            self.assertFalse(
                any(
                    record.get("artifact_type") == "pdf"
                    for record in report_data["document_verification"]
                )
            )

    def test_create_document_from_cache_removes_stale_pdf_verification_when_conversion_fails(self):
        song_list = [{"Artist": "The Campfire Trio", "Title": "Trail Song"}]
        lyrics_cache = {"The Campfire Trio - Trail Song": "First line\nSecond line"}

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "lyrics.docx"
            document_quality_path = tmp_path / "data" / "review" / "document_quality.json"

            def fake_convert(docx_path, pdf_path=None, runner=None, which=None):  # noqa: ARG001
                target_path = Path(docx_path).with_suffix(".pdf")
                write_minimal_pdf(target_path, "x" * 80)
                return target_path, None

            with patch("app.document_creation.convert_document_to_pdf", side_effect=fake_convert):
                self._run_create_document(
                    tmp_path,
                    song_list,
                    lyrics_cache,
                    {},
                    lyrics_output=output_path,
                    pdf_output=True,
                )

            state, errors = load_document_verification(document_quality_path)
            self.assertEqual(errors, [])
            self.assertIn(str(output_path.with_suffix(".pdf")), state["entries"])

            with patch(
                "app.document_creation.convert_document_to_pdf",
                return_value=(None, "converter missing"),
            ):
                report_data = self._run_create_document(
                    tmp_path,
                    song_list,
                    lyrics_cache,
                    {},
                    lyrics_output=output_path,
                    pdf_output=True,
                )

            self.assertEqual(len(report_data["pdf_errors"]), 1)
            self.assertEqual(report_data["pdf_errors"][0]["reason"], "converter missing")

            state, errors = load_document_verification(document_quality_path)
            self.assertEqual(errors, [])
            self.assertNotIn(str(output_path.with_suffix(".pdf")), state["entries"])

    def test_create_document_from_cache_persists_document_verification_for_generated_artifacts(self):
        song_list = [{"Artist": "The Campfire Trio", "Title": "Trail Song"}]
        lyrics_cache = {"The Campfire Trio - Trail Song": "First line\nSecond line"}

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "lyrics.docx"
            document_quality_path = tmp_path / "data" / "review" / "document_quality.json"

            report_data = self._run_create_document(
                tmp_path,
                song_list,
                lyrics_cache,
                {},
                lyrics_output=output_path,
            )

            state, errors = load_document_verification(document_quality_path)

            self.assertEqual(errors, [])
            self.assertEqual(len(report_data["document_verification"]), 2)
            self.assertEqual(len(report_data["review_gate_decisions"]), 2)
            self.assertIn(str(output_path), state["entries"])
            self.assertIn(str(output_path.with_suffix(".md")), state["entries"])
            self.assertEqual(
                state["entries"][str(output_path)]["verification_status"],
                "passed",
            )
            self.assertTrue(report_data["review_gate_decisions"][0]["review_ready"])
            self.assertEqual(
                report_data["review_gate_decisions"][0]["failure_reasons"],
                [],
            )
            self.assertEqual(
                state["entries"][str(output_path.with_suffix(".md"))]["verification_reasons"],
                ["meets_neatness_thresholds"],
            )

    def test_create_document_from_cache_flags_sparse_generated_artifacts(self):
        song_list = [
            {"Artist": "The Campfire Trio", "Title": "Trail Song"},
            {"Artist": "The Campfire Trio", "Title": "River Song"},
        ]
        lyrics_cache = {
            "The Campfire Trio - Trail Song": "Only line",
            "The Campfire Trio - River Song": "Solo",
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            output_path = tmp_path / "lyrics.docx"
            document_quality_path = tmp_path / "data" / "review" / "document_quality.json"

            self._run_create_document(
                tmp_path,
                song_list,
                lyrics_cache,
                {},
                lyrics_output=output_path,
            )

            state, errors = load_document_verification(document_quality_path)

            self.assertEqual(errors, [])
            self.assertEqual(
                state["entries"][str(output_path)]["verification_status"],
                "failed",
            )
            self.assertIn(
                "fragmented_song_blocks",
                state["entries"][str(output_path)]["verification_reasons"],
            )
            self.assertIn(
                "sparse_layout",
                state["entries"][str(output_path.with_suffix(".md"))]["verification_reasons"],
            )


if __name__ == "__main__":
    unittest.main()
