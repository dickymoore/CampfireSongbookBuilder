import unittest

from docx.oxml.ns import qn

from tests.docx_stub import install_docx_stub

install_docx_stub()

from docx import Document  # noqa: E402

from app.document_formatting import (  # noqa: E402
    add_contents_page,
    create_two_column_section,
    set_document_mirrored_margins,
)


class TestDocumentFormatting(unittest.TestCase):
    def test_create_two_column_section_reuses_existing_section(self):
        document = Document()

        self.assertEqual(len(document.sections), 1)

        create_two_column_section(document)

        self.assertEqual(len(document.sections), 1)
        cols = document.sections[0]._sectPr.xpath('./w:cols')[0]
        attrs = getattr(cols, 'attrs', None)
        if attrs is not None:
            self.assertEqual(attrs.get('w:num'), '2')
            self.assertEqual(attrs.get('w:space'), '720')
        else:
            self.assertEqual(cols.get(qn('w:num')), '2')
            self.assertEqual(cols.get(qn('w:space')), '720')

    def test_add_contents_page_builds_a_table_with_page_references(self):
        document = Document()

        add_contents_page(
            document,
            [
                {
                    "Artist": "The Campfire Trio",
                    "Title": "River Song",
                    "bookmark_name": "song_1_the_campfire_trio_river_song",
                },
                {
                    "Artist": "The Campfire Trio",
                    "Title": "Trail Song",
                    "bookmark_name": "song_2_the_campfire_trio_trail_song",
                },
            ],
            generated_at="2026-07-09T17:50:45+01:00",
        )

        self.assertEqual(document.paragraphs[0].text, "2 Campfire Songs | 2026-07-09")
        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual(table.cell(0, 0).text, "Song")
        self.assertEqual(table.cell(0, 1).text, "Page")
        self.assertEqual(table.cell(1, 0).text, "River Song - The Campfire Trio")
        self.assertEqual(table.cell(2, 0).text, "Trail Song - The Campfire Trio")
        if hasattr(table.cell(1, 1).paragraphs[0], "_p"):
            self.assertIn("PAGEREF song_1_the_campfire_trio_river_song", table.cell(1, 1).paragraphs[0]._p.xml)
            self.assertIn("PAGEREF song_2_the_campfire_trio_trail_song", table.cell(2, 1).paragraphs[0]._p.xml)
        if hasattr(table, "_tbl"):
            self.assertIn("tblBorders", table._tbl.xml)
            self.assertIn("tblLayout", table._tbl.xml)
        if hasattr(table.cell(0, 0), "_tc") and hasattr(table.cell(0, 0)._tc, "xml"):
            self.assertIn("noWrap", table.cell(0, 0)._tc.xml)

    def test_set_document_mirrored_margins_enables_mirror_margins_and_sets_inside_outside(self):
        document = Document()

        set_document_mirrored_margins(document, 0.79, 0.01)

        section = document.sections[0]
        left_inches = getattr(section.left_margin, "inches", section.left_margin)
        right_inches = getattr(section.right_margin, "inches", section.right_margin)
        gutter_inches = getattr(section.gutter, "inches", section.gutter)
        self.assertAlmostEqual(left_inches, 0.01, places=2)
        self.assertAlmostEqual(right_inches, 0.01, places=2)
        self.assertAlmostEqual(gutter_inches, 0.78, places=2)
        if hasattr(document, "settings"):
            self.assertIsNotNone(document.settings._element.find(qn("w:mirrorMargins")))
