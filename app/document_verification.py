import json
import logging
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
except ModuleNotFoundError:  # pragma: no cover - exercised in test environments without python-docx
    from tests.docx_stub import install_docx_stub

    install_docx_stub()
    from docx import Document


logger = logging.getLogger(__name__)

DOCUMENT_VERIFICATION_VERSION = 1
DEFAULT_DOCUMENT_QUALITY_PATH = Path("data/review/document_quality.json")
ARTIFACT_TYPES = ("markdown", "docx", "pdf")
VERIFICATION_STATUSES = ("passed", "failed")
NEATNESS_REASON_PASSED = "meets_neatness_thresholds"
NEATNESS_REASON_EXCESSIVE_WHITESPACE = "excessive_whitespace"
NEATNESS_REASON_SPARSE_LAYOUT = "sparse_layout"
NEATNESS_REASON_FRAGMENTED_SONG_BLOCKS = "fragmented_song_blocks"
NEATNESS_REASON_PRINT_HOSTILE_STRUCTURE = "print_hostile_structure"
MAX_CONSECUTIVE_BLANK_LINES = 2
MAX_BLANK_LINE_RATIO = 0.35
MIN_AVERAGE_BODY_LINES_PER_SONG = 2.0
SHORT_SONG_BLOCK_MAX_LINES = 1
MIN_FRAGMENTED_BLOCKS = 2
MIN_FRAGMENTED_BLOCK_RATIO = 0.5
PDF_REASON_MISSING = "missing_pdf"
PDF_REASON_EMPTY = "empty_pdf"


def _validation_error(file_path, field, reason):
    return {
        "file_path": str(file_path),
        "field": field,
        "reason": reason,
    }


def _empty_document_verification_state(updated_at=None):
    return {
        "version": DOCUMENT_VERIFICATION_VERSION,
        "updated_at": updated_at,
        "entries": {},
    }


def _now_iso():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _require_text(field_name, value):
    if not isinstance(value, str) or value == "":
        raise ValueError("{} must be a non-empty string; got {!r}".format(field_name, value))
    return value


def _validate_enum(field_name, value, allowed_values):
    if value not in allowed_values:
        raise ValueError(
            "{} must be one of {}; got {!r}".format(
                field_name,
                ", ".join(repr(item) for item in allowed_values),
                value,
            )
        )
    return value


def validate_artifact_type(artifact_type):
    return _validate_enum("artifact_type", artifact_type, ARTIFACT_TYPES)


def validate_verification_status(verification_status):
    return _validate_enum(
        "verification_status",
        verification_status,
        VERIFICATION_STATUSES,
    )


def build_document_verification_record(
    artifact_path,
    artifact_type,
    verification_status,
    verification_reasons=None,
    verified_at=None,
):
    artifact_path_value = _require_text("artifact_path", artifact_path)
    artifact_type_value = validate_artifact_type(artifact_type)
    verification_status_value = validate_verification_status(verification_status)

    if verification_reasons is None:
        verification_reason_items = []
    elif not isinstance(verification_reasons, list):
        raise ValueError(
            "verification_reasons must be a list; got {!r}".format(
                verification_reasons
            )
        )
    else:
        verification_reason_items = verification_reasons

    reason_values = []
    for index, reason in enumerate(verification_reason_items):
        reason_values.append(
            _require_text("verification_reasons[{}]".format(index), reason)
        )

    verified_at_value = _require_text(
        "verified_at",
        verified_at if verified_at is not None else _now_iso(),
    )

    return {
        "artifact_path": artifact_path_value,
        "artifact_type": artifact_type_value,
        "verification_status": verification_status_value,
        "verification_reasons": reason_values,
        "verified_at": verified_at_value,
    }


def _count_max_consecutive_blank_lines(lines):
    max_run = 0
    current_run = 0
    for line in lines:
        if isinstance(line, str) and line.strip() == "":
            current_run += 1
            if current_run > max_run:
                max_run = current_run
        else:
            current_run = 0
    return max_run


def _blank_line_ratio(lines):
    if not lines:
        return 0.0
    blank_lines = sum(1 for line in lines if isinstance(line, str) and line.strip() == "")
    return blank_lines / float(len(lines))


def _song_block_body_lines(song_blocks):
    return [len(block.get("body_lines", [])) for block in song_blocks]


def _evaluate_neatness_reasons(all_lines, song_blocks):
    reasons = []
    body_line_counts = _song_block_body_lines(song_blocks)

    if not song_blocks or sum(body_line_counts) == 0:
        reasons.append(NEATNESS_REASON_PRINT_HOSTILE_STRUCTURE)

    if (
        _count_max_consecutive_blank_lines(all_lines) > MAX_CONSECUTIVE_BLANK_LINES
        or _blank_line_ratio(all_lines) > MAX_BLANK_LINE_RATIO
    ):
        reasons.append(NEATNESS_REASON_EXCESSIVE_WHITESPACE)

    if song_blocks:
        average_body_lines = sum(body_line_counts) / float(len(song_blocks))
        if average_body_lines < MIN_AVERAGE_BODY_LINES_PER_SONG:
            reasons.append(NEATNESS_REASON_SPARSE_LAYOUT)

        short_block_count = sum(
            1 for count in body_line_counts if count <= SHORT_SONG_BLOCK_MAX_LINES
        )
        if (
            short_block_count >= MIN_FRAGMENTED_BLOCKS
            and short_block_count / float(len(song_blocks)) >= MIN_FRAGMENTED_BLOCK_RATIO
        ):
            reasons.append(NEATNESS_REASON_FRAGMENTED_SONG_BLOCKS)

    if not reasons:
        return [NEATNESS_REASON_PASSED]
    return reasons


def _markdown_song_blocks(markdown_text):
    lines = markdown_text.splitlines()
    song_blocks = []
    current_block = None
    in_code_block = False

    for line in lines:
        if line.startswith("# "):
            current_block = {"heading": line, "body_lines": []}
            song_blocks.append(current_block)
            in_code_block = False
            continue

        if current_block is None:
            continue

        if line.startswith("```"):
            in_code_block = not in_code_block
            continue

        if in_code_block:
            if line.strip() != "":
                current_block["body_lines"].append(line)
            continue

        if line.strip() != "":
            current_block["body_lines"].append(line)

    return lines, song_blocks


def _docx_song_blocks(document):
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    song_blocks = []
    current_block = None

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text
        style_name = getattr(getattr(paragraph, "style", None), "name", "")
        is_heading = isinstance(style_name, str) and style_name.startswith("Heading")

        if is_heading and paragraph_text.strip() != "":
            current_block = {"heading": paragraph_text, "body_lines": []}
            song_blocks.append(current_block)
            continue

        if current_block is None:
            continue

        for line in paragraph_text.splitlines():
            if line.strip() != "":
                current_block["body_lines"].append(line)

    return paragraphs, song_blocks


def evaluate_markdown_artifact_neatness(markdown_text):
    all_lines, song_blocks = _markdown_song_blocks(markdown_text)
    reasons = _evaluate_neatness_reasons(all_lines, song_blocks)
    return {
        "verification_status": "passed" if reasons == [NEATNESS_REASON_PASSED] else "failed",
        "verification_reasons": reasons,
    }


def evaluate_docx_artifact_neatness(document):
    all_lines, song_blocks = _docx_song_blocks(document)
    reasons = _evaluate_neatness_reasons(all_lines, song_blocks)
    return {
        "verification_status": "passed" if reasons == [NEATNESS_REASON_PASSED] else "failed",
        "verification_reasons": reasons,
    }


def evaluate_pdf_artifact_neatness(pdf_path):
    pdf_path = Path(pdf_path)
    if not pdf_path.exists():
        return {
            "verification_status": "failed",
            "verification_reasons": [PDF_REASON_MISSING],
        }

    if pdf_path.stat().st_size <= 0:
        return {
            "verification_status": "failed",
            "verification_reasons": [PDF_REASON_EMPTY],
        }

    return {
        "verification_status": "passed",
        "verification_reasons": [NEATNESS_REASON_PASSED],
    }


def evaluate_document_artifact(artifact_path, artifact_type=None, verified_at=None):
    artifact_path = Path(artifact_path)
    if artifact_type is None:
        extension = artifact_path.suffix.lstrip(".").lower()
        artifact_type = "markdown" if extension == "md" else extension
    elif isinstance(artifact_type, str):
        artifact_type = artifact_type.lower()

    artifact_type_value = validate_artifact_type(artifact_type)

    if artifact_type_value == "markdown":
        evaluation = evaluate_markdown_artifact_neatness(
            artifact_path.read_text(encoding="utf-8")
        )
    elif artifact_type_value == "docx":
        evaluation = evaluate_docx_artifact_neatness(Document(artifact_path))
    else:
        evaluation = evaluate_pdf_artifact_neatness(artifact_path)

    return build_document_verification_record(
        artifact_path=str(artifact_path),
        artifact_type=artifact_type_value,
        verification_status=evaluation["verification_status"],
        verification_reasons=evaluation["verification_reasons"],
        verified_at=verified_at,
    )


def _load_json_document(file_path):
    if not file_path.exists():
        return None, []

    try:
        with file_path.open("r", encoding="utf-8") as handle:
            return json.load(handle), []
    except json.JSONDecodeError as exc:
        return None, [
            _validation_error(
                file_path,
                "$",
                "Invalid JSON in document verification file: {}".format(exc),
            )
        ]
    except OSError as exc:
        return None, [
            _validation_error(
                file_path,
                "$",
                "Failed to read document verification file: {}".format(exc),
            )
        ]


def _validate_required_text(file_path, field_path, value, errors):
    if not isinstance(value, str) or value == "":
        errors.append(
            _validation_error(
                file_path,
                field_path,
                "{} must be a non-empty string; got {!r}".format(
                    field_path.rsplit(".", 1)[-1],
                    value,
                ),
            )
        )
        return None
    return value


def _validate_verification_reasons(file_path, field_path, value, errors):
    if value is None:
        errors.append(
            _validation_error(
                file_path,
                field_path,
                "verification_reasons must be a list; got None",
            )
        )
        return None

    if not isinstance(value, list):
        errors.append(
            _validation_error(
                file_path,
                field_path,
                "verification_reasons must be a list; got {!r}".format(value),
            )
        )
        return None

    validated_reasons = []
    for index, reason in enumerate(value):
        validated_reason = _validate_required_text(
            file_path,
            "{}[{}]".format(field_path, index),
            reason,
            errors,
        )
        if validated_reason is not None:
            validated_reasons.append(validated_reason)

    return validated_reasons


def _validate_document_verification_record(file_path, outer_artifact_path, record, errors):
    record_field_base = "entries['{}']".format(outer_artifact_path)

    if not isinstance(record, dict):
        errors.append(
            _validation_error(
                file_path,
                record_field_base,
                "document verification entries must be dictionaries; got {!r}".format(record),
            )
        )
        return None

    artifact_path = _validate_required_text(
        file_path,
        "{}.artifact_path".format(record_field_base),
        record.get("artifact_path"),
        errors,
    )

    artifact_type = record.get("artifact_type")
    if artifact_type is None:
        errors.append(
            _validation_error(
                file_path,
                "{}.artifact_type".format(record_field_base),
                "artifact_type must be a non-empty string; got None",
            )
        )
        validated_artifact_type = None
    else:
        try:
            validated_artifact_type = validate_artifact_type(artifact_type)
        except ValueError as exc:
            errors.append(
                _validation_error(
                    file_path,
                    "{}.artifact_type".format(record_field_base),
                    str(exc),
                )
            )
            validated_artifact_type = None

    verification_status = record.get("verification_status")
    if verification_status is None:
        errors.append(
            _validation_error(
                file_path,
                "{}.verification_status".format(record_field_base),
                "verification_status must be a non-empty string; got None",
            )
        )
        validated_verification_status = None
    else:
        try:
            validated_verification_status = validate_verification_status(
                verification_status
            )
        except ValueError as exc:
            errors.append(
                _validation_error(
                    file_path,
                    "{}.verification_status".format(record_field_base),
                    str(exc),
                )
            )
            validated_verification_status = None

    validated_reasons = _validate_verification_reasons(
        file_path,
        "{}.verification_reasons".format(record_field_base),
        record.get("verification_reasons"),
        errors,
    )
    verified_at = _validate_required_text(
        file_path,
        "{}.verified_at".format(record_field_base),
        record.get("verified_at"),
        errors,
    )

    if (
        artifact_path is None
        or validated_artifact_type is None
        or validated_verification_status is None
        or validated_reasons is None
        or verified_at is None
    ):
        return None

    if artifact_path != outer_artifact_path:
        errors.append(
            _validation_error(
                file_path,
                "{}.artifact_path".format(record_field_base),
                "entry key must match artifact_path {!r}; got {!r}".format(
                    artifact_path,
                    outer_artifact_path,
                ),
            )
        )
        return None

    return build_document_verification_record(
        artifact_path=artifact_path,
        artifact_type=validated_artifact_type,
        verification_status=validated_verification_status,
        verification_reasons=validated_reasons,
        verified_at=verified_at,
    )


def load_document_verification(file_path=DEFAULT_DOCUMENT_QUALITY_PATH):
    file_path = Path(file_path)
    document, errors = _load_json_document(file_path)

    if document is None:
        if not file_path.exists():
            return _empty_document_verification_state(), []
        return _empty_document_verification_state(), errors

    if not isinstance(document, dict):
        return _empty_document_verification_state(), [
            _validation_error(
                file_path,
                "$",
                "document verification file must contain a top-level object; got {!r}".format(
                    document
                ),
            )
        ]

    state_errors = []
    version = document.get("version", DOCUMENT_VERIFICATION_VERSION)
    if version != DOCUMENT_VERIFICATION_VERSION:
        state_errors.append(
            _validation_error(
                file_path,
                "version",
                "version must be {}; got {!r}".format(
                    DOCUMENT_VERIFICATION_VERSION,
                    version,
                ),
            )
        )
        version = DOCUMENT_VERIFICATION_VERSION

    updated_at = document.get("updated_at")
    if updated_at is not None and not isinstance(updated_at, str):
        state_errors.append(
            _validation_error(
                file_path,
                "updated_at",
                "updated_at must be a string or null; got {!r}".format(updated_at),
            )
        )
        updated_at = None

    entries = document.get("entries")
    if entries is None:
        return _empty_document_verification_state(updated_at=updated_at), (
            errors
            + state_errors
            + [
                _validation_error(
                    file_path,
                    "entries",
                    "entries is required and must be an object keyed by artifact_path; got None",
                )
            ]
        )

    if not isinstance(entries, dict):
        return _empty_document_verification_state(updated_at=updated_at), (
            errors
            + state_errors
            + [
                _validation_error(
                    file_path,
                    "entries",
                    "entries must be an object keyed by artifact_path; got {!r}".format(
                        entries
                    ),
                )
            ]
        )

    validated_entries = {}
    for outer_artifact_path, record in entries.items():
        validated_record = _validate_document_verification_record(
            file_path,
            outer_artifact_path,
            record,
            state_errors,
        )
        if validated_record is not None:
            validated_entries[outer_artifact_path] = validated_record

    return {
        "version": version,
        "updated_at": updated_at,
        "entries": validated_entries,
    }, errors + state_errors


def build_document_verification_state(document_verification_records, updated_at=None):
    entries = {}
    for document_verification in document_verification_records:
        if not isinstance(document_verification, dict):
            raise ValueError(
                "document verification records must be dictionaries; got {!r}".format(
                    document_verification
                )
            )

        record = build_document_verification_record(
            artifact_path=document_verification.get("artifact_path"),
            artifact_type=document_verification.get("artifact_type"),
            verification_status=document_verification.get("verification_status"),
            verification_reasons=document_verification.get("verification_reasons"),
            verified_at=document_verification.get("verified_at"),
        )
        entries[record["artifact_path"]] = record

    return {
        "version": DOCUMENT_VERIFICATION_VERSION,
        "updated_at": updated_at if updated_at is not None else _now_iso(),
        "entries": entries,
    }


def save_document_verification(file_path, document_verification_records, updated_at=None):
    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)

    state = build_document_verification_state(
        document_verification_records,
        updated_at=updated_at,
    )
    with file_path.open("w", encoding="utf-8") as handle:
        json.dump(state, handle, ensure_ascii=False, indent=2, sort_keys=True)
        handle.write("\n")

    logger.info("Saved document verification state to %s", file_path)
    return state
